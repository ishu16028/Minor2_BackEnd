# app/services/email_service.py
import os
import base64
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from flask import current_app
import docx
import re

# Gmail API scopes
SCOPES = ['https://www.googleapis.com/auth/gmail.compose']


def get_gmail_service(token_info):
    """Create a Gmail API service using user's token information"""
    credentials = Credentials.from_authorized_user_info(info=token_info, scopes=SCOPES)

    # Refresh token if expired
    if credentials and credentials.expired and credentials.refresh_token:
        credentials.refresh(Request())

    service = build('gmail', 'v1', credentials=credentials)
    return service


def parse_docx_template(template_path):
    """Extract content from a DOCX template file"""
    doc = docx.Document(template_path)
    content = "\n".join([para.text for para in doc.paragraphs])
    return content


def personalize_template(template_content, user_data, resume_data, company_name=None, hr_name=None):
    """Replace placeholders in template with personalized information"""
    # Map of placeholders to actual data
    replacements = {
        "{{user_name}}": resume_data.get("name", user_data.get("name", "")),
        "{{user_email}}": resume_data.get("contact_details", {}).get("email", user_data.get("email", "")),
        "{{user_phone}}": resume_data.get("contact_details", {}).get("phone", ""),
        "{{linkedin_url}}": resume_data.get("social_links", {}).get("LinkedIn", ""),
        "{{github_url}}": resume_data.get("social_links", {}).get("GitHub", ""),
        "{{skills}}": ", ".join(resume_data.get("skills", [])),
        "{{company_name}}": company_name or "[COMPANY NAME]",
        "{{hr_name}}": hr_name or "[HR NAME]",
        "{{current_date}}": "[CURRENT DATE]"  # You might want to replace this with actual date
    }

    # Replace placeholders
    personalized_content = template_content
    for placeholder, value in replacements.items():
        personalized_content = personalized_content.replace(placeholder, value)

    return personalized_content


def create_email_draft(service, to, subject, body, user_id='me'):
    """Create an email draft in user's Gmail account"""
    message = MIMEMultipart()
    message['to'] = to
    message['subject'] = subject

    msg = MIMEText(body)
    message.attach(msg)

    raw = base64.urlsafe_b64encode(message.as_bytes()).decode()

    try:
        draft = service.users().drafts().create(
            userId=user_id,
            body={'message': {'raw': raw}}
        ).execute()
        return {'success': True, 'draft_id': draft['id']}
    except Exception as e:
        current_app.logger.error(f"Error creating email draft: {str(e)}")
        return {'success': False, 'error': str(e)}


def get_available_templates():
    """Get list of available email templates from the templates directory"""
    templates_dir = os.path.join(current_app.static_folder, 'email_templates')
    templates = []

    if os.path.exists(templates_dir):
        for filename in os.listdir(templates_dir):
            if filename.endswith('.docx'):
                template_path = os.path.join(templates_dir, filename)
                template_name = os.path.splitext(filename)[0].replace('_', ' ').title()
                templates.append({
                    'id': filename,
                    'name': template_name,
                    'path': template_path
                })

    return templates


def prepare_cold_email(template_id, user_id, hr_contact_id, company_name=None, custom_subject=None):
    """Prepare cold email content based on template and user data"""
    from app.models.user import User
    from app.models.resume import Resume
    from app.models.hr_contact import HRContact

    # Get user data
    user = User.query.get(user_id)
    if not user:
        return {'success': False, 'error': 'User not found'}

    # Get resume data
    resume = Resume.query.filter_by(user_id=user_id).first()
    if not resume:
        return {'success': False, 'error': 'Resume not found'}

    # Get HR contact data if provided
    hr_contact = None
    if hr_contact_id:
        hr_contact = HRContact.query.get(hr_contact_id)

    # Get template
    templates = get_available_templates()
    template = next((t for t in templates if t['id'] == template_id), None)
    if not template:
        return {'success': False, 'error': 'Email template not found'}

    # Parse template
    template_content = parse_docx_template(template['path'])

    # Personalize template
    hr_name = hr_contact.contact_name if hr_contact else None
    company = company_name or (hr_contact.company_name if hr_contact else None)
    personalized_content = personalize_template(
        template_content,
        user.__dict__,
        resume.data,
        company,
        hr_name
    )

    # Generate subject
    if not custom_subject:
        # Extract first line as subject or use default
        subject_match = re.search(r'^.*$', personalized_content, re.MULTILINE)
        subject = subject_match.group(0) if subject_match else f"Job Application - {resume.data.get('name', '')}"
    else:
        subject = custom_subject

    return {
        'success': True,
        'to': hr_contact.email if hr_contact else '',
        'subject': subject,
        'body': personalized_content
    }